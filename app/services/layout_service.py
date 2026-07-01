import base64
import hashlib
import io
import json
import os
import platform
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from typing import List

from fastapi import HTTPException
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation


BASE_DIR = Path(__file__).resolve().parents[2]
STATIC_DIR = BASE_DIR / "app" / "static" / "layout_picker"
CACHE_DIR = STATIC_DIR / "preview_cache"
TEMPLATE_PATH = BASE_DIR / "ppt template export" / "merkle_template.potx"

POTX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument."
    "presentationml.template.main+xml"
)
PPTX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument."
    "presentationml.presentation.main+xml"
)

PREVIEW_WIDTH = 640
PREVIEW_HEIGHT = 360
RENDERED_PREVIEW_CACHE = None
RENDERED_PREVIEW_SOURCE = None
LAST_RENDER_ERROR = None

def _clean_hex_color(value):
    color = str(value or "").strip()

    if color.startswith("#"):
        color = color[1:]

    if len(color) != 6:
        return None

    try:
        int(color, 16)

    except Exception:
        return None

    return f"#{color.upper()}"


def _is_black_or_white(color):
    normalized = _clean_hex_color(color)

    return normalized in {
        "#000000",
        "#FFFFFF",
    }


def _extract_template_palette():
    if not TEMPLATE_PATH.exists():
        return []

    color_counts = {}
    source_lookup = {}

    try:
        with zipfile.ZipFile(TEMPLATE_PATH, "r") as package:
            for name in package.namelist():
                if not name.endswith(".xml"):
                    continue

                if not (
                    name.startswith("ppt/theme/")
                    or name.startswith("ppt/slideMasters/")
                    or name.startswith("ppt/slideLayouts/")
                    or name.startswith("ppt/slides/")
                ):
                    continue

                text = package.read(name).decode("utf-8", errors="ignore")

                matches = re.findall(
                    r"<a:srgbClr[^>]*\sval=\"([0-9A-Fa-f]{6})\"",
                    text
                )

                matches += re.findall(
                    r"\s(?:color|fill|stroke)=\"?#?([0-9A-Fa-f]{6})\"?",
                    text
                )

                for match in matches:
                    color = _clean_hex_color(match)

                    if not color:
                        continue

                    if _is_black_or_white(color):
                        continue

                    color_counts[color] = color_counts.get(color, 0) + 1
                    source_lookup.setdefault(color, set()).add(name)

    except Exception:
        return []

    return [
        {
            "hex": color,
            "count": color_counts[color],
            "sources": sorted(source_lookup.get(color, []))[:5],
        }
        for color in sorted(
            color_counts,
            key=lambda item: color_counts[item],
            reverse=True
        )
    ]


def _font(size=14):
    try:
        return ImageFont.truetype("arial.ttf", size)
    except Exception:
        return ImageFont.load_default()


def _load_template_presentation():
    if not TEMPLATE_PATH.exists():
        raise HTTPException(
            status_code=404,
            detail=f"Template not found: {TEMPLATE_PATH}"
        )

    converted = tempfile.NamedTemporaryFile(delete=False, suffix=".pptx")
    converted.close()

    with zipfile.ZipFile(TEMPLATE_PATH, "r") as source:
        with zipfile.ZipFile(converted.name, "w", zipfile.ZIP_DEFLATED) as target:
            for item in source.infolist():
                data = source.read(item.filename)

                if item.filename == "[Content_Types].xml":
                    data = (
                        data.decode("utf-8")
                        .replace(POTX_CONTENT_TYPE, PPTX_CONTENT_TYPE)
                        .encode("utf-8")
                    )

                target.writestr(item, data)

    return Presentation(converted.name)


def _find_libreoffice():
    configured_path = os.environ.get("LIBREOFFICE_PATH")

    if configured_path and Path(configured_path).exists():
        return configured_path

    for command in ("soffice", "libreoffice"):
        resolved = shutil.which(command)

        if resolved:
            return resolved

    windows_candidates = [
        Path("C:/Program Files/LibreOffice/program/soffice.exe"),
        Path("C:/Program Files (x86)/LibreOffice/program/soffice.exe"),
    ]

    for candidate in windows_candidates:
        if candidate.exists():
            return str(candidate)

    return None


def _powerpoint_available():
    if platform.system().lower() != "windows":
        return False

    try:
        import win32com.client  # noqa: F401
        return True

    except Exception:
        return False


def _relative_luminance(rgb):
    red, green, blue = rgb[:3]
    return (
        0.2126 * red
        + 0.7152 * green
        + 0.0722 * blue
    )


def _outline_color_for_box(image, box):
    x1, y1, x2, y2 = box
    sample_points = [
        (
            max(x1 + 4, 0),
            max(y1 + 4, 0)
        ),
        (
            min(x2 - 4, image.width - 1),
            max(y1 + 4, 0)
        ),
        (
            max(x1 + 4, 0),
            min(y2 - 4, image.height - 1)
        ),
        (
            min(x2 - 4, image.width - 1),
            min(y2 - 4, image.height - 1)
        ),
        (
            min(max(int((x1 + x2) / 2), 0), image.width - 1),
            min(max(int((y1 + y2) / 2), 0), image.height - 1)
        ),
    ]

    luminance = sum(
        _relative_luminance(
            image.getpixel(point)
        )
        for point in sample_points
    ) / len(sample_points)

    if luminance < 120:
        return "#FFFFFF", "#111827"

    return "#111827", "#FFFFFF"


def _draw_placeholder_outline(draw, box, label, line_color, label_fill):
    x1, y1, x2, y2 = box
    dash = 8
    gap = 5

    for x in range(x1, x2, dash + gap):
        draw.line(
            (
                x,
                y1,
                min(x + dash, x2),
                y1
            ),
            fill=line_color,
            width=2
        )
        draw.line(
            (
                x,
                y2,
                min(x + dash, x2),
                y2
            ),
            fill=line_color,
            width=2
        )

    for y in range(y1, y2, dash + gap):
        draw.line(
            (
                x1,
                y,
                x1,
                min(y + dash, y2)
            ),
            fill=line_color,
            width=2
        )
        draw.line(
            (
                x2,
                y,
                x2,
                min(y + dash, y2)
            ),
            fill=line_color,
            width=2
        )

    if (x2 - x1) > 90 and (y2 - y1) > 28:
        draw.rectangle(
            (
                x1 + 5,
                y1 + 5,
                min(x1 + 135, x2 - 5),
                y1 + 24
            ),
            fill=label_fill
        )
        draw.text(
            (
                x1 + 9,
                y1 + 7
            ),
            label[:22],
            fill=line_color,
            font=_font(10)
        )


def _draw_icon_box(draw, box, line_color, label):
    x1, y1, x2, y2 = box

    draw.rounded_rectangle(
        box,
        radius=3,
        outline=line_color,
        width=1
    )

    cx = int(
        (x1 + x2) / 2
    )
    cy = int(
        (y1 + y2) / 2
    )

    if label == "table":
        for offset in (-5, 0, 5):
            draw.line(
                (
                    x1 + 4,
                    cy + offset,
                    x2 - 4,
                    cy + offset
                ),
                fill=line_color,
                width=1
            )
            draw.line(
                (
                    cx + offset,
                    y1 + 4,
                    cx + offset,
                    y2 - 4
                ),
                fill=line_color,
                width=1
            )

    elif label == "chart":
        for idx, height in enumerate((7, 12, 17)):
            left = x1 + 5 + (idx * 7)
            draw.rectangle(
                (
                    left,
                    y2 - 5 - height,
                    left + 4,
                    y2 - 5
                ),
                fill=line_color
            )

    elif label == "image":
        draw.rectangle(
            (
                x1 + 5,
                y1 + 5,
                x2 - 5,
                y2 - 5
            ),
            outline=line_color,
            width=1
        )
        draw.polygon(
            [
                (
                    x1 + 6,
                    y2 - 6
                ),
                (
                    cx - 2,
                    cy
                ),
                (
                    cx + 3,
                    cy + 5
                ),
                (
                    x2 - 6,
                    y2 - 6
                ),
            ],
            outline=line_color
        )

    elif label == "video":
        draw.rectangle(
            (
                x1 + 5,
                y1 + 6,
                x2 - 5,
                y2 - 6
            ),
            outline=line_color,
            width=1
        )
        draw.polygon(
            [
                (
                    cx - 3,
                    cy - 6
                ),
                (
                    cx - 3,
                    cy + 6
                ),
                (
                    cx + 7,
                    cy
                ),
            ],
            fill=line_color
        )

    else:
        draw.ellipse(
            (
                cx - 7,
                cy - 7,
                cx + 7,
                cy + 7
            ),
            outline=line_color,
            width=1
        )


def _draw_placeholder_insert_icons(draw, box, kind, line_color):
    x1, y1, x2, y2 = box
    width = x2 - x1
    height = y2 - y1

    if width < 110 or height < 70:
        return

    if kind == "Picture":
        icons = ["image"]
    elif kind == "Content":
        icons = [
            "table",
            "chart",
            "image",
            "video",
        ]
    else:
        return

    icon_size = 28
    gap = 7
    total_width = (len(icons) * icon_size) + ((len(icons) - 1) * gap)
    start_x = int(
        x1 + ((width - total_width) / 2)
    )
    start_y = int(
        y1 + ((height - icon_size) / 2)
    )

    for index, icon in enumerate(icons):
        icon_left = start_x + (index * (icon_size + gap))
        _draw_icon_box(
            draw,
            (
                icon_left,
                start_y,
                icon_left + icon_size,
                start_y + icon_size
            ),
            line_color,
            icon
        )


def _overlay_placeholder_outlines(canvas, prs, layout):
    draw = ImageDraw.Draw(
        canvas
    )

    for placeholder in layout.placeholders:
        kind = _placeholder_kind(
            placeholder
        )
        box = _scaled_box(
            placeholder,
            prs.slide_width,
            prs.slide_height
        )
        line_color, label_fill = _outline_color_for_box(
            canvas,
            box
        )
        _draw_placeholder_outline(
            draw,
            box,
            kind,
            line_color,
            label_fill
        )
        _draw_placeholder_insert_icons(
            draw,
            box,
            kind,
            line_color
        )


def _image_file_to_data_url(path, prs=None, layout=None):
    with Image.open(path) as image:
        return _image_to_data_url(
            image,
            prs,
            layout
        )


def _image_to_data_url(image, prs=None, layout=None):
    image = image.copy()
    image.thumbnail(
        (
            PREVIEW_WIDTH,
            PREVIEW_HEIGHT
        )
    )

    canvas = Image.new(
        "RGB",
        (
            PREVIEW_WIDTH,
            PREVIEW_HEIGHT
        ),
        "#F8FAFC"
    )

    left = int(
        (PREVIEW_WIDTH - image.width) / 2
    )
    top = int(
        (PREVIEW_HEIGHT - image.height) / 2
    )

    canvas.paste(
        image.convert("RGB"),
        (
            left,
            top
        )
    )

    if prs is not None and layout is not None:
        _overlay_placeholder_outlines(
            canvas,
            prs,
            layout
        )

    buffer = io.BytesIO()
    canvas.save(
        buffer,
        format="PNG"
    )

    encoded = base64.b64encode(
        buffer.getvalue()
    ).decode("ascii")

    return f"data:image/png;base64,{encoded}"


def _create_layout_preview_deck(prs):
    preview_prs = _load_template_presentation()

    for index in range(len(preview_prs.slide_layouts)):
        preview_prs.slides.add_slide(
            preview_prs.slide_layouts[index]
        )

    temp_file = tempfile.NamedTemporaryFile(
        delete=False,
        suffix=".pptx"
    )

    preview_prs.save(
        temp_file.name
    )

    return temp_file.name


def _preview_cache_key():

    if not TEMPLATE_PATH.exists():
        return None

    stat = TEMPLATE_PATH.stat()
    raw_key = (
        f"{TEMPLATE_PATH.resolve()}|"
        f"{stat.st_mtime_ns}|"
        f"{stat.st_size}"
    )

    return hashlib.sha256(
        raw_key.encode("utf-8")
    ).hexdigest()[:24]


def _load_preview_cache(layout_count):

    cache_key = _preview_cache_key()

    if not cache_key:
        return None

    cache_file = CACHE_DIR / f"{cache_key}.json"

    if not cache_file.exists():
        return None

    try:
        payload = json.loads(
            cache_file.read_text(
                encoding="utf-8"
            )
        )

        previews = payload.get(
            "previews"
        )

        if (
            isinstance(previews, list)
            and len(previews) == layout_count
        ):
            return payload

    except Exception:
        return None

    return None


def _save_preview_cache(mode, renderer, previews):

    cache_key = _preview_cache_key()

    if not cache_key or not previews:
        return

    try:
        CACHE_DIR.mkdir(
            parents=True,
            exist_ok=True
        )

        cache_file = CACHE_DIR / f"{cache_key}.json"

        cache_file.write_text(
            json.dumps({
                "mode": mode,
                "renderer": renderer,
                "previews": previews,
            }),
            encoding="utf-8"
        )

    except Exception:
        return


def _render_layout_previews_with_powerpoint(prs):
    global LAST_RENDER_ERROR

    if not _powerpoint_available():
        return None

    try:
        import pythoncom
        import win32com.client

    except Exception:
        return None

    deck_path = str(
        Path(_create_layout_preview_deck(prs)).resolve()
    )

    output_dir = Path(
        tempfile.mkdtemp(
            prefix="layout_poc_powerpoint_previews_"
        )
    )

    powerpoint = None
    presentation = None

    try:
        pythoncom.CoInitialize()

        powerpoint = win32com.client.DispatchEx(
            "PowerPoint.Application"
        )

        powerpoint.Visible = 1

        presentation = powerpoint.Presentations.Open(
            deck_path,
            WithWindow=False
        )

        previews = []

        for index in range(1, presentation.Slides.Count + 1):
            image_path = output_dir / f"layout_{index:03d}.png"

            presentation.Slides(index).Export(
                str(image_path),
                "PNG",
                PREVIEW_WIDTH,
                PREVIEW_HEIGHT
            )

            previews.append(
                _image_file_to_data_url(
                    image_path,
                    prs,
                    prs.slide_layouts[index - 1]
                )
            )

        if len(previews) < len(prs.slide_layouts):
            return None

        return previews[:len(prs.slide_layouts)]

    except Exception as exc:
        LAST_RENDER_ERROR = f"PowerPoint renderer failed: {type(exc).__name__}: {exc}"
        return None

    finally:
        try:
            if presentation:
                presentation.Close()

        except Exception:
            pass

        try:
            if powerpoint:
                powerpoint.Quit()

        except Exception:
            pass

        try:
            pythoncom.CoUninitialize()

        except Exception:
            pass

        try:
            Path(deck_path).unlink(
                missing_ok=True
            )

        except Exception:
            pass

        try:
            shutil.rmtree(
                output_dir,
                ignore_errors=True
            )

        except Exception:
            pass


def _render_layout_previews_with_libreoffice(prs):
    global LAST_RENDER_ERROR

    libreoffice = _find_libreoffice()

    if not libreoffice:
        LAST_RENDER_ERROR = "LibreOffice renderer failed: soffice/libreoffice command not found"
        return None

    deck_path = _create_layout_preview_deck(
        prs
    )

    output_dir = Path(
        tempfile.mkdtemp(
            prefix="layout_poc_previews_"
        )
    )
    profile_dir = Path(
        tempfile.mkdtemp(
            prefix="layout_poc_libreoffice_profile_"
        )
    )

    try:
        subprocess.run(
            [
                libreoffice,
                "--headless",
                "--nologo",
                "--nodefault",
                "--nofirststartwizard",
                "--nolockcheck",
                f"-env:UserInstallation=file://{profile_dir}",
                "--convert-to",
                "pdf",
                "--outdir",
                str(output_dir),
                deck_path,
            ],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=90,
        )

        pdf_files = sorted(
            list(output_dir.glob("*.pdf"))
            + list(output_dir.glob("*.PDF"))
        )

        if not pdf_files:
            LAST_RENDER_ERROR = "LibreOffice renderer failed: no PDF was produced"
            return None

        previews = _render_pdf_pages_to_data_urls(
            pdf_files[0],
            prs
        )

        if len(previews) < len(prs.slide_layouts):
            LAST_RENDER_ERROR = (
                "LibreOffice renderer failed: rendered fewer pages "
                f"({len(previews)}) than layouts ({len(prs.slide_layouts)})"
            )
            return None

        return [
            _image_to_data_url(
                previews[index],
                prs,
                prs.slide_layouts[index]
            )
            for index in range(len(prs.slide_layouts))
        ]

    except Exception as exc:
        stderr = ""

        if isinstance(exc, subprocess.CalledProcessError):
            stderr = (
                exc.stderr or b""
            ).decode(
                "utf-8",
                errors="ignore"
            ).strip()

        detail = f": {stderr}" if stderr else ""
        LAST_RENDER_ERROR = f"LibreOffice renderer failed: {type(exc).__name__}: {exc}{detail}"
        return None

    finally:
        try:
            Path(deck_path).unlink(
                missing_ok=True
            )

        except Exception:
            pass

        try:
            shutil.rmtree(
                output_dir,
                ignore_errors=True
            )

        except Exception:
            pass

        try:
            shutil.rmtree(
                profile_dir,
                ignore_errors=True
            )

        except Exception:
            pass


def _render_pdf_pages_to_data_urls(pdf_path, prs):
    global LAST_RENDER_ERROR

    try:
        import fitz

    except Exception as exc:
        LAST_RENDER_ERROR = (
            "LibreOffice renderer failed: PyMuPDF is required to convert "
            f"PDF pages to PNG previews ({type(exc).__name__}: {exc})"
        )
        return []

    pages = []

    with fitz.open(pdf_path) as document:
        for page in document:
            pixmap = page.get_pixmap(
                matrix=fitz.Matrix(2, 2),
                alpha=False
            )
            image = Image.open(
                io.BytesIO(
                    pixmap.tobytes("png")
                )
            ).convert("RGB")
            pages.append(
                image
            )

    return pages


def _render_layout_previews(prs):
    global RENDERED_PREVIEW_CACHE
    global RENDERED_PREVIEW_SOURCE
    global LAST_RENDER_ERROR

    LAST_RENDER_ERROR = None

    source_key = (
        str(TEMPLATE_PATH),
        TEMPLATE_PATH.stat().st_mtime
    )

    if (
        RENDERED_PREVIEW_CACHE is not None
        and RENDERED_PREVIEW_SOURCE == source_key
    ):
        return RENDERED_PREVIEW_CACHE

    cached = _load_preview_cache(
        len(prs.slide_layouts)
    )

    if cached:

        RENDERED_PREVIEW_CACHE = {
            "mode": cached.get("mode", "cached"),
            "renderer": cached.get("renderer"),
            "previews": cached.get("previews"),
        }
        RENDERED_PREVIEW_SOURCE = source_key
        return RENDERED_PREVIEW_CACHE

    renderers = [
        (
            "powerpoint",
            "Microsoft PowerPoint",
            _render_layout_previews_with_powerpoint,
        ),
        (
            "libreoffice",
            _find_libreoffice() or "LibreOffice",
            _render_layout_previews_with_libreoffice,
        ),
    ]

    for mode, renderer, render_func in renderers:
        previews = render_func(
            prs
        )

        if previews:
            _save_preview_cache(
                mode,
                renderer,
                previews
            )

            RENDERED_PREVIEW_CACHE = {
                "mode": mode,
                "renderer": renderer,
                "previews": previews,
            }
            RENDERED_PREVIEW_SOURCE = source_key
            return RENDERED_PREVIEW_CACHE

    return {
        "mode": "schematic",
        "renderer": None,
        "previews": None,
        "error": LAST_RENDER_ERROR,
    }


def _placeholder_kind(placeholder):
    try:
        placeholder_type = str(placeholder.placeholder_format.type)
    except Exception:
        placeholder_type = ""

    if "TITLE" in placeholder_type:
        return "Title"
    if "PICTURE" in placeholder_type:
        return "Picture"
    if "BODY" in placeholder_type:
        return "Body"
    if "OBJECT" in placeholder_type:
        return "Content"
    return "Placeholder"


def _placeholder_colors(kind):
    colors = {
        "Title": ("#DDEBFF", "#3B73B9"),
        "Body": ("#E9F7EF", "#39895A"),
        "Content": ("#FFF3D6", "#B87800"),
        "Picture": ("#F2E8FF", "#8056B3"),
        "Placeholder": ("#EEF2F7", "#64748B"),
    }
    return colors.get(kind, colors["Placeholder"])


def _scaled_box(shape, slide_width, slide_height):
    left = int(shape.left / slide_width * PREVIEW_WIDTH)
    top = int(shape.top / slide_height * PREVIEW_HEIGHT)
    width = int(shape.width / slide_width * PREVIEW_WIDTH)
    height = int(shape.height / slide_height * PREVIEW_HEIGHT)

    return (
        left,
        top,
        max(left + width, left + 1),
        max(top + height, top + 1),
    )


def _layout_preview_data_url(prs, layout):
    image = Image.new("RGB", (PREVIEW_WIDTH, PREVIEW_HEIGHT), "#F8FAFC")
    draw = ImageDraw.Draw(image)

    draw.rectangle(
        (0, 0, PREVIEW_WIDTH - 1, PREVIEW_HEIGHT - 1),
        outline="#CBD5E1",
        width=2,
    )
    draw.text((16, 14), layout.name, fill="#0F172A", font=_font(18))

    for placeholder in layout.placeholders:
        kind = _placeholder_kind(placeholder)
        fill, outline = _placeholder_colors(kind)
        box = _scaled_box(placeholder, prs.slide_width, prs.slide_height)
        draw.rounded_rectangle(box, radius=8, fill=fill, outline=outline, width=2)
        draw.text(
            (box[0] + 8, box[1] + 7),
            f"{kind}: {placeholder.name}"[:38],
            fill=outline,
            font=_font(12),
        )

    if not layout.placeholders:
        draw.text((24, 76), "No placeholders", fill="#64748B", font=_font(15))

    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    encoded = base64.b64encode(buffer.getvalue()).decode("ascii")
    return f"data:image/png;base64,{encoded}"


def _layout_metadata(prs, layout, index, rendered_preview=None, preview_mode=None):
    placeholders = []

    for placeholder in layout.placeholders:
        placeholder_format = placeholder.placeholder_format
        placeholders.append({
            "idx": placeholder_format.idx,
            "type": str(placeholder_format.type),
            "kind": _placeholder_kind(placeholder),
            "name": placeholder.name,
            "left": int(placeholder.left),
            "top": int(placeholder.top),
            "width": int(placeholder.width),
            "height": int(placeholder.height),
        })

    placeholder_kinds = {
        item["kind"]
        for item in placeholders
    }

    return {
        "index": index,
        "name": layout.name,
        "placeholder_count": len(placeholders),
        "placeholders": placeholders,
        "has_content_placeholder": (
            "Content" in placeholder_kinds
            or "Body" in placeholder_kinds
        ),
        "preview": rendered_preview or _layout_preview_data_url(prs, layout),
        "preview_mode": preview_mode if rendered_preview else "schematic",
    }


def _delete_preview_cache():
    cache_key = _preview_cache_key()

    if not cache_key:
        return

    cache_file = CACHE_DIR / f"{cache_key}.json"

    try:
        if cache_file.exists():
            cache_file.unlink()

    except Exception:
        return


def get_layouts(template_path=None, force_refresh=False):
    global TEMPLATE_PATH
    global RENDERED_PREVIEW_CACHE
    global RENDERED_PREVIEW_SOURCE

    original_template_path = TEMPLATE_PATH

    if template_path:
        TEMPLATE_PATH = Path(template_path)

    try:
        if force_refresh:
            RENDERED_PREVIEW_CACHE = None
            RENDERED_PREVIEW_SOURCE = None
            _delete_preview_cache()

        prs = _load_template_presentation()
        preview_result = _render_layout_previews(
            prs
        )
        previews = preview_result.get(
            "previews"
        )

        return {
            "template": str(TEMPLATE_PATH),
            "preview_mode": preview_result.get(
                "mode"
            ),
            "renderer": preview_result.get(
                "renderer"
            ),
            "renderer_error": preview_result.get(
                "error"
            ),
            "palette": _extract_template_palette(),
            "slide_width": int(prs.slide_width),
            "slide_height": int(prs.slide_height),
            "layouts": [
                _layout_metadata(
                    prs,
                    layout,
                    index,
                    (
                        previews[index]
                        if previews
                        else None
                    ),
                    preview_result.get("mode")
                )
                for index, layout in enumerate(prs.slide_layouts)
            ],
        }

    finally:
        TEMPLATE_PATH = original_template_path
