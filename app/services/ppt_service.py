# services/ppt_service.py

import base64
import io
import json
import tempfile
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.chart import XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER


TEMPLATE_PATH = (
    Path(__file__).resolve().parents[2]
    / "ppt template export"
    / "merkle_template.potx"
)


PPT_CONFIG_PATH = (
    Path(__file__).resolve().parents[1]
    / "config"
    / "chart_colors.json"
)


POTX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument."
    "presentationml.template.main+xml"
)


PPTX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument."
    "presentationml.presentation.main+xml"
)


def _load_base_presentation():

    if not TEMPLATE_PATH.exists():
        return Presentation()

    converted = tempfile.NamedTemporaryFile(
        delete=False,
        suffix=".pptx"
    )

    converted.close()

    with zipfile.ZipFile(
        TEMPLATE_PATH,
        "r"
    ) as source:

        with zipfile.ZipFile(
            converted.name,
            "w",
            zipfile.ZIP_DEFLATED
        ) as target:

            for item in source.infolist():

                data = source.read(
                    item.filename
                )

                if item.filename == "[Content_Types].xml":

                    text = data.decode(
                        "utf-8"
                    )

                    text = text.replace(
                        POTX_CONTENT_TYPE,
                        PPTX_CONTENT_TYPE
                    )

                    data = text.encode(
                        "utf-8"
                    )

                target.writestr(
                    item,
                    data
                )

    try:
        return Presentation(
            converted.name
        )

    except Exception:
        return Presentation()


def _inches_from_emu(value):

    return value / 914400


def _set_title_slide_text(
    slide,
    title,
    subtitle
):

    title_set = False
    subtitle_set = False

    for shape in slide.placeholders:

        placeholder_type = shape.placeholder_format.type

        if (
            not title_set
            and placeholder_type == PP_PLACEHOLDER.TITLE
        ):

            shape.text = title
            title_set = True
            continue

        if (
            not subtitle_set
            and placeholder_type in [
                PP_PLACEHOLDER.SUBTITLE,
                PP_PLACEHOLDER.BODY,
            ]
        ):

            shape.text = subtitle
            subtitle_set = True

    if not title_set:

        box = slide.shapes.add_textbox(
            Inches(0.75),
            Inches(2.3),
            Inches(11.8),
            Inches(0.8)
        )

        box.text_frame.text = title
        box.text_frame.paragraphs[0].font.size = Pt(34)

    if not subtitle_set:

        box = slide.shapes.add_textbox(
            Inches(0.75),
            Inches(3.25),
            Inches(11.8),
            Inches(0.5)
        )

        box.text_frame.text = subtitle
        box.text_frame.paragraphs[0].font.size = Pt(16)


def _find_slide_layout(
    prs,
    preferred_names,
    fallback_index
):

    preferred = [
        name.lower()
        for name in preferred_names
    ]

    for layout in prs.slide_layouts:

        layout_name = layout.name.lower()

        if layout_name in preferred:
            return layout

    for layout in prs.slide_layouts:

        layout_name = layout.name.lower()

        if any(
            name in layout_name
            for name in preferred
        ):
            return layout

    return prs.slide_layouts[
        min(
            fallback_index,
            len(prs.slide_layouts) - 1
        )
    ]


def _merge_ppt_config(base_config, override_config):

    if not isinstance(override_config, dict):
        return base_config

    merged_config = dict(base_config)

    palette = override_config.get(
        "palette"
    )

    if isinstance(palette, list) and palette:
        merged_config["palette"] = palette

    for key in (
        "title",
        "subtitle",
        "background_color",
        "title_color",
        "text_color",
    ):

        if key in override_config:
            merged_config[key] = override_config.get(
                key
            )

    for key in (
        "header",
        "footer",
        "logo",
    ):

        configured_values = override_config.get(
            key
        )

        if isinstance(configured_values, dict):

            merged_values = dict(
                merged_config.get(
                    key,
                    {}
                )
            )

            merged_values.update(
                configured_values
            )

            merged_config[key] = merged_values

    return merged_config


def _load_ppt_config(override_config=None):

    default_config = {
        "palette": [
            "#4C78A8",
            "#F58518",
            "#54A24B",
            "#E45756",
            "#72B7B2",
            "#EECA3B",
        ],
        "title": "Analysis Summary",
        "subtitle": "Generated by PPT Export Service",
        "background_color": "#FFFFFF",
        "title_color": "#1F2933",
        "text_color": "#1F2933",
        "header": {
            "text": "",
            "height": 0.35,
            "background_color": "#FFFFFF",
            "text_color": "#1F2933",
        },
        "footer": {
            "text": "",
            "show_slide_number": True,
            "text_color": "#5B6770",
        },
        "logo": {
            "path": "",
            "base64": "",
            "width": 1.0,
            "height": 0.35,
            "left": 0.45,
            "top": 0.12,
        },
    }

    if not PPT_CONFIG_PATH.exists():
        return _merge_ppt_config(
            default_config,
            override_config
        )

    try:

        with PPT_CONFIG_PATH.open(
            "r",
            encoding="utf-8"
        ) as file:

            loaded = json.load(file)

        if not isinstance(loaded, dict):
            return _merge_ppt_config(
                default_config,
                override_config
            )

        default_config = _merge_ppt_config(
            default_config,
            loaded
        )

        return _merge_ppt_config(
            default_config,
            override_config
        )

    except Exception:
        return _merge_ppt_config(
            default_config,
            override_config
        )


def _normalize_color_key(value):

    return str(value or "").strip().lower()


def _clean_hex_color(value):

    color = str(value or "").strip()

    if color.startswith("#"):
        color = color[1:]

    if len(color) != 6:
        return None

    try:
        int(color, 16)

    except Exception:
        return None

    return color.upper()


def _color_lookup(mapping, key):

    normalized_key = _normalize_color_key(key)

    for candidate, color in mapping.items():

        if _normalize_color_key(candidate) == normalized_key:
            return _clean_hex_color(color)

    return None


def _palette_color(config, index):

    palette = config.get(
        "palette",
        []
    )

    if not palette:
        return None

    return _clean_hex_color(
        palette[index % len(palette)]
    )


def _apply_fill_color(target, hex_color):

    color = _clean_hex_color(
        hex_color
    )

    if not color:
        return

    try:

        target.format.fill.solid()
        target.format.fill.fore_color.rgb = (
            RGBColor.from_string(color)
        )

    except Exception:
        return


def _set_font_color(paragraph, hex_color):

    color = _clean_hex_color(
        hex_color
    )

    if not color:
        return

    try:
        paragraph.font.color.rgb = RGBColor.from_string(
            color
        )

    except Exception:
        return


def _set_slide_background(slide, hex_color):

    color = _clean_hex_color(
        hex_color
    )

    if not color:
        return

    try:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor.from_string(
            color
        )

    except Exception:
        return


def _logo_source(logo_config):

    if not isinstance(logo_config, dict):
        return None

    base64_value = str(
        logo_config.get(
            "base64",
            ""
        )
        or ""
    ).strip()

    if base64_value:

        if "," in base64_value:
            base64_value = base64_value.split(
                ",",
                1
            )[1]

        try:
            return io.BytesIO(
                base64.b64decode(
                    base64_value
                )
            )

        except Exception:
            return None

    logo_path = str(
        logo_config.get(
            "path",
            ""
        )
        or ""
    ).strip()

    if logo_path and Path(logo_path).exists():
        return logo_path

    return None


def _apply_fallback_branding(
    slide,
    slide_width,
    slide_height,
    ppt_config,
    slide_number=None
):

    _set_slide_background(
        slide,
        ppt_config.get(
            "background_color"
        )
    )

    header = ppt_config.get(
        "header",
        {}
    )

    footer = ppt_config.get(
        "footer",
        {}
    )

    logo = ppt_config.get(
        "logo",
        {}
    )

    header_height = float(
        header.get(
            "height",
            0.35
        )
        or 0
    )

    if header_height > 0:

        header_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0),
            Inches(0),
            Inches(slide_width),
            Inches(header_height)
        )

        _apply_fill_color(
            header_shape,
            header.get(
                "background_color"
            )
        )

        header_shape.line.fill.background()

        header_text = str(
            header.get(
                "text",
                ""
            )
            or ""
        ).strip()

        if header_text:

            header_box = slide.shapes.add_textbox(
                Inches(1.6),
                Inches(0.08),
                Inches(slide_width - 2.1),
                Inches(max(header_height - 0.08, 0.2))
            )

            paragraph = header_box.text_frame.paragraphs[0]
            paragraph.text = header_text
            paragraph.font.size = Pt(9)
            _set_font_color(
                paragraph,
                header.get(
                    "text_color"
                )
            )

    logo_file = _logo_source(
        logo
    )

    if logo_file:

        try:
            slide.shapes.add_picture(
                logo_file,
                Inches(float(logo.get("left", 0.45))),
                Inches(float(logo.get("top", 0.12))),
                width=Inches(float(logo.get("width", 1.0))),
                height=Inches(float(logo.get("height", 0.35)))
            )

        except Exception:
            pass

    footer_text = str(
        footer.get(
            "text",
            ""
        )
        or ""
    ).strip()

    if (
        footer.get(
            "show_slide_number",
            True
        )
        and slide_number is not None
    ):

        page_text = str(
            slide_number
        )

        footer_text = (
            f"{footer_text} | {page_text}"
            if footer_text
            else page_text
        )

    if footer_text:

        footer_box = slide.shapes.add_textbox(
            Inches(0.5),
            Inches(slide_height - 0.35),
            Inches(slide_width - 1.0),
            Inches(0.2)
        )

        paragraph = footer_box.text_frame.paragraphs[0]
        paragraph.text = footer_text
        paragraph.font.size = Pt(8)
        _set_font_color(
            paragraph,
            footer.get(
                "text_color"
            )
        )


def _apply_chart_colors(
    chart,
    parsed_chart,
    ppt_config
):

    series_items = parsed_chart.get(
        "series"
    )

    if series_items:

        for idx, series in enumerate(chart.series):

            color = _palette_color(
                ppt_config,
                idx
            )

            _apply_fill_color(
                series,
                color
            )

        return

    if not chart.series:
        return

    first_series = chart.series[0]

    for idx, point in enumerate(first_series.points):

        color = _palette_color(
            ppt_config,
            idx
        )

        _apply_fill_color(
            point,
            color
        )


def _chart_complexity(parsed_chart):

    if not parsed_chart:
        return 0

    category_count = len(
        parsed_chart.get(
            "categories",
            []
        )
    )

    series_count = len(
        parsed_chart.get(
            "series",
            []
        )
    ) or 1

    return category_count * series_count


def _content_layout(
    slide_width,
    slide_height,
    bullet_count,
    parsed_chart=None
):

    margin_x = 0.5
    title_top = 0.3
    title_height = 0.6
    content_top = 1.05
    bottom_margin = 0.55
    gutter = 0.28

    content_height = (
        slide_height
        - content_top
        - bottom_margin
    )

    has_chart = parsed_chart is not None

    if not has_chart:

        return {
            "title": (
                margin_x,
                title_top,
                slide_width - (margin_x * 2),
                title_height,
            ),
            "body": (
                margin_x,
                content_top,
                slide_width - (margin_x * 2),
                content_height,
            ),
            "chart": None,
            "body_font": 16,
        }

    complexity = _chart_complexity(
        parsed_chart
    )

    if bullet_count <= 1:
        body_width = 2.15
    elif bullet_count == 2:
        body_width = 2.55
    else:
        body_width = 2.95

    if complexity >= 40:
        body_width = max(
            2.05,
            body_width - 0.35
        )
    elif complexity <= 8:
        body_width = min(
            3.2,
            body_width + 0.25
        )

    chart_left = margin_x + body_width + gutter
    chart_width = slide_width - chart_left - margin_x

    if chart_width < 6.0:

        chart_width = 6.0
        body_width = max(
            1.8,
            slide_width
            - margin_x
            - gutter
            - chart_width
            - margin_x
        )
        chart_left = margin_x + body_width + gutter

    body_font = 12

    if bullet_count <= 1:
        body_font = 13
    elif bullet_count >= 3:
        body_font = 11

    return {
        "title": (
            margin_x,
            title_top,
            slide_width - (margin_x * 2),
            title_height,
        ),
        "body": (
            margin_x,
            content_top + 0.15,
            body_width,
            content_height - 0.15,
        ),
        "chart": (
            chart_left,
            content_top,
            chart_width,
            content_height,
        ),
        "body_font": body_font,
    }


# ============================================
# CHART PARSER
# ============================================

def _apply_vega_transforms(chart_spec, data_values):

    rows = list(data_values)

    for transform in chart_spec.get(
        "transform",
        []
    ):

        fold_fields = transform.get(
            "fold"
        )

        as_fields = transform.get(
            "as",
            []
        )

        if (
            not fold_fields
            or len(as_fields) < 2
        ):
            continue

        key_field = as_fields[0]
        value_field = as_fields[1]
        folded_rows = []

        for row in rows:

            for field in fold_fields:

                folded = dict(row)
                folded[key_field] = field
                folded[value_field] = row.get(
                    field,
                    0
                )
                folded_rows.append(folded)

        rows = folded_rows

    return rows


def _fallback_field(row, candidates):

    for field in candidates:

        if field in row and row.get(field) is not None:
            return field

    return None


def _to_float(value):

    try:
        return float(value)

    except Exception:
        return 0


def extract_chart_data(chart_spec):

    if not chart_spec:
        return None

    # ============================================
    # STRING -> JSON
    # ============================================

    if isinstance(chart_spec, str):

        try:
            chart_spec = json.loads(chart_spec)

        except Exception:
            return None

    # ============================================
    # SIMPLE FORMAT
    # ============================================

    if (
        "labels" in chart_spec
        and "values" in chart_spec
    ):

        return {

            "chart_type": chart_spec.get(
                "type",
                "bar"
            ),

            "orientation": "vertical",

            "title": chart_spec.get(
                "title",
                "Chart"
            ),

            "categories": chart_spec["labels"],

            "values": chart_spec["values"],
        }

    # ============================================
    # VEGA-LITE FORMAT
    # ============================================

    try:

        mark = chart_spec.get(
            "mark",
            "bar"
        )

        if isinstance(mark, dict):

            mark = mark.get(
                "type",
                "bar"
            )

        data_values = (
            chart_spec
            .get("data", {})
            .get("values", [])
        )

        data_values = _apply_vega_transforms(
            chart_spec,
            data_values
        )

        encoding = chart_spec.get(
            "encoding",
            {}
        )

        x_encoding = encoding.get(
            "x",
            {}
        )

        y_encoding = encoding.get(
            "y",
            {}
        )

        x_field = x_encoding.get(
            "field"
        )

        y_field = y_encoding.get(
            "field"
        )

        x_type = x_encoding.get(
            "type"
        )

        y_type = y_encoding.get(
            "type"
        )

        color_field = (
            encoding
            .get("color", {})
            .get("field")
        )

        y_offset_field = (
            encoding
            .get("yOffset", {})
            .get("field")
        )

        x_offset_field = (
            encoding
            .get("xOffset", {})
            .get("field")
        )

        series_field = (
            y_offset_field
            or x_offset_field
            or color_field
        )

        if (
            series_field
            and data_values
            and series_field not in data_values[0]
        ):

            series_field = _fallback_field(
                data_values[0],
                [
                    "METRIC_KEY",
                    "ROLE_KEY",
                    "REG_KEY",
                    "IND_KEY",
                    "REGION",
                    "INDUSTRY",
                ]
            )

        categories = []
        values = []

        # ============================================
        # DETERMINE ORIENTATION
        # ============================================

        if mark == "rect":

            orientation = "horizontal"

            category_field = y_field or _fallback_field(
                data_values[0],
                [
                    "ATTRIBUTE",
                    "BRAND",
                    "QID",
                    "QID_LABEL",
                ]
            )

            value_field = _fallback_field(
                data_values[0],
                [
                    "VALUE",
                    "PCT",
                    "RECORD_COUNT",
                    "RESPONDENT_COUNT",
                ]
            )

            series_field = x_field or series_field

            if (
                series_field
                and data_values
                and series_field not in data_values[0]
            ):

                series_field = _fallback_field(
                    data_values[0],
                    [
                        "REG_KEY",
                        "IND_KEY",
                        "METRIC_KEY",
                        "ROLE_KEY",
                        "REGION",
                        "INDUSTRY",
                    ]
                )

        elif (
            x_type == "quantitative"
            and y_type == "nominal"
        ):

            orientation = "horizontal"

            category_field = y_field
            value_field = x_field

        else:

            orientation = "vertical"

            category_field = x_field
            value_field = y_field

        if not category_field:

            category_field = _fallback_field(
                data_values[0],
                [
                    "QID",
                    "QID_LABEL",
                    "US_REGION",
                    "BRAND",
                    "ATTRIBUTE",
                ]
            )

        if not value_field:

            value_field = _fallback_field(
                data_values[0],
                [
                    "VALUE",
                    "PCT",
                    "RECORD_COUNT",
                    "RESPONDENT_COUNT",
                ]
            )

        # ============================================
        # EXTRACT VALUES
        # ============================================

        if series_field and series_field in data_values[0]:

            categories = []
            category_seen = set()
            series_names = []
            series_seen = set()
            value_lookup = {}

            for row in data_values:

                category = str(
                    row.get(
                        category_field
                    )
                )

                series_name = str(
                    row.get(
                        series_field
                    )
                )

                if category not in category_seen:
                    categories.append(category)
                    category_seen.add(category)

                if series_name not in series_seen:
                    series_names.append(series_name)
                    series_seen.add(series_name)

                value_lookup[
                    (
                        series_name,
                        category
                    )
                ] = _to_float(
                    row.get(
                        value_field,
                        0
                    )
                )

            series = []

            for series_name in series_names:

                series.append({
                    "name": series_name,
                    "values": [
                        value_lookup.get(
                            (
                                series_name,
                                category
                            ),
                            0
                        )
                        for category in categories
                    ]
                })

            return {

                "chart_type": (
                    "bar"
                    if mark == "rect"
                    else mark
                ),

                "orientation": orientation,

                "title": chart_spec.get(
                    "title",
                    "Chart"
                ),

                "categories": categories,

                "series": series,
            }

        for row in data_values:

            categories.append(
                str(
                    row.get(
                        category_field
                    )
                )
            )

            try:

                values.append(
                    _to_float(
                        row.get(
                            value_field,
                            0
                        )
                    )
                )

            except Exception:

                values.append(0)

        return {

                "chart_type": (
                    "bar"
                    if mark == "rect"
                    else mark
                ),

            "orientation": orientation,

            "title": chart_spec.get(
                "title",
                "Chart"
            ),

            "categories": categories,

            "values": values,
        }

    except Exception as e:

        print(
            "Chart parsing error:",
            str(e)
        )

        return None


# ============================================
# CHART TYPE
# ============================================

def get_chart_type(parsed_chart):

    chart_type = parsed_chart.get(
        "chart_type",
        "bar"
    )

    orientation = parsed_chart.get(
        "orientation",
        "vertical"
    )

    # ============================================
    # LINE
    # ============================================

    if chart_type == "line":

        return XL_CHART_TYPE.LINE

    # ============================================
    # PIE / DONUT
    # ============================================

    if chart_type in ["pie", "arc"]:

        return XL_CHART_TYPE.PIE

    # ============================================
    # HORIZONTAL BAR
    # ============================================

    if orientation == "horizontal":

        return XL_CHART_TYPE.BAR_CLUSTERED

    # ============================================
    # VERTICAL BAR
    # ============================================

    return XL_CHART_TYPE.COLUMN_CLUSTERED


# ============================================
# PPT GENERATOR
# ============================================

def generate_pptx(
    presentation_json: dict,
    include_charts: bool = True,
    ppt_config=None
):

    template_available = TEMPLATE_PATH.exists()

    prs = _load_base_presentation()

    if not template_available:
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

    slide_width = _inches_from_emu(
        prs.slide_width
    )

    slide_height = _inches_from_emu(
        prs.slide_height
    )

    scale_x = slide_width / 13.333
    scale_y = slide_height / 7.5

    fallback_config = None

    if not template_available:

        fallback_config = _load_ppt_config(
            ppt_config
        )

    title = presentation_json.get(
        "presentation_title",
        "Analysis Summary"
    )

    subtitle = "Generated by PPT Export Service"

    if fallback_config:

        title = fallback_config.get(
            "title",
            title
        ) or title

        subtitle = fallback_config.get(
            "subtitle",
            subtitle
        ) or subtitle

    # ============================================
    # TITLE SLIDE
    # ============================================

    title_slide_layout = _find_slide_layout(
        prs,
        [
            "Title Page",
            "Title Slide"
        ],
        0
    )

    title_slide = prs.slides.add_slide(
        title_slide_layout
    )

    _set_title_slide_text(
        slide=title_slide,
        title=title,
        subtitle=subtitle
    )

    if fallback_config:

        _apply_fallback_branding(
            slide=title_slide,
            slide_width=slide_width,
            slide_height=slide_height,
            ppt_config=fallback_config,
            slide_number=1
        )

    # ============================================
    # LOOP THROUGH SLIDES
    # ============================================

    for slide_data in presentation_json.get(
        "slides",
        []
    ):

        slide_layout = _find_slide_layout(
            prs,
            [
                "Custom Layout",
                "Blank"
            ],
            6
        )

        slide = prs.slides.add_slide(
            slide_layout
        )

        slide_number = len(prs.slides)

        if fallback_config:

            _apply_fallback_branding(
                slide=slide,
                slide_width=slide_width,
                slide_height=slide_height,
                ppt_config=fallback_config,
                slide_number=slide_number
            )

        parsed_chart = None

        if include_charts:

            parsed_chart = extract_chart_data(
                slide_data.get(
                    "chart_spec"
                )
            )

        bullets = slide_data["slide"].get(
            "bullets",
            []
        )

        layout = _content_layout(
            slide_width=slide_width,
            slide_height=slide_height,
            bullet_count=len(bullets),
            parsed_chart=parsed_chart
        )

        # ============================================
        # TITLE
        # ============================================

        title_left, title_top, title_width, title_height = (
            layout["title"]
        )

        title_box = slide.shapes.add_textbox(
            Inches(title_left),
            Inches(title_top),
            Inches(title_width),
            Inches(title_height)
        )

        title_tf = title_box.text_frame

        title_tf.text = slide_data["slide"].get(
            "title",
            "Key Insights"
        )

        title_tf.paragraphs[0].font.size = Pt(22)

        if fallback_config:

            _set_font_color(
                title_tf.paragraphs[0],
                fallback_config.get(
                    "title_color"
                )
            )

        # ============================================
        # BULLETS
        # ============================================

        body_left, body_top, body_width, body_height = (
            layout["body"]
        )

        body_box = slide.shapes.add_textbox(
            Inches(body_left),
            Inches(body_top),
            Inches(body_width),
            Inches(body_height)
        )

        tf = body_box.text_frame

        tf.clear()

        tf.word_wrap = True

        for idx, bullet in enumerate(bullets):

            if idx == 0:

                p = tf.paragraphs[0]

            else:

                p = tf.add_paragraph()

            p.text = f"- {bullet}"

            p.font.size = Pt(
                layout["body_font"]
            )

            p.space_after = Pt(7)

            if fallback_config:

                _set_font_color(
                    p,
                    fallback_config.get(
                        "text_color"
                    )
                )

        # ============================================
        # CHART
        # ============================================

        if parsed_chart and layout["chart"]:

                chart_data = CategoryChartData()

                chart_data.categories = (
                    parsed_chart["categories"]
                )

                series = parsed_chart.get(
                    "series"
                )

                if series:

                    for item in series:

                        chart_data.add_series(
                            item.get(
                                "name",
                                "Series"
                            ),
                            item.get(
                                "values",
                                []
                            )
                        )

                else:

                    chart_data.add_series(
                        "Series 1",
                        parsed_chart["values"]
                    )

                chart_left, chart_top, chart_width, chart_height = (
                    layout["chart"]
                )

                chart = slide.shapes.add_chart(

                    get_chart_type(
                        parsed_chart
                    ),

                    Inches(chart_left),

                    Inches(chart_top),

                    Inches(chart_width),

                    Inches(chart_height),

                    chart_data

                ).chart

                chart.has_title = True

                chart.chart_title.text_frame.text = (
                    parsed_chart.get(
                        "title",
                        "Chart"
                    )
                )

                chart.has_legend = bool(series)

                if fallback_config:

                    _apply_chart_colors(
                        chart=chart,
                        parsed_chart=parsed_chart,
                        ppt_config=fallback_config
                    )

                if chart.has_legend:

                    chart.legend.position = (
                        XL_LEGEND_POSITION.BOTTOM
                    )

                    chart.legend.include_in_layout = False

                try:

                    category_font_size = 6 if (
                        len(parsed_chart["categories"]) > 12
                    ) else 8

                    chart.category_axis.tick_labels.font.size = Pt(
                        category_font_size
                    )
                    chart.value_axis.tick_labels.font.size = Pt(8)

                except Exception:

                    pass

    # ============================================
    # SAVE
    # ============================================

    temp_file = tempfile.NamedTemporaryFile(
        delete=False,
        suffix=".pptx"
    )

    prs.save(temp_file.name)

    return temp_file.name
