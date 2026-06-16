import base64
import io
import os
import platform
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from typing import List

from fastapi import FastAPI, HTTPException
from fastapi.responses import FileResponse, HTMLResponse
from fastapi.staticfiles import StaticFiles
from PIL import Image, ImageDraw, ImageFont
from pydantic import BaseModel
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Inches, Pt


BASE_DIR = Path(__file__).resolve().parents[1]
STATIC_DIR = Path(__file__).resolve().parent / "static"
TEMPLATE_PATH = BASE_DIR / "ppt template export" / "merkle_template.potx"

POTX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument."
    "presentationml.template.main+xml"
)
PPTX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument."
    "presentationml.presentation.main+xml"
)

PREVIEW_WIDTH = 640
PREVIEW_HEIGHT = 360
RENDERED_PREVIEW_CACHE = None
RENDERED_PREVIEW_SOURCE = None
LAST_RENDER_ERROR = None

app = FastAPI(title="Template Layout POC")
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")


class RenderSlideRequest(BaseModel):
    layout_index: int
    title: str = "Sample Insights"
    bullets: List[str] = [
        "Familiarity increased to 61%.",
        "Consideration is strongest in the Midwest.",
        "Usage is trailing awareness by 19 points.",
    ]
    chart_title: str = "Sample Chart"
    categories: List[str] = ["Familiarity", "Consideration", "Usage"]
    values: List[float] = [61, 42, 28]
    include_chart: bool = True
    selected_colors: List[str] = []


def _clean_hex_color(value):
    color = str(value or "").strip()

    if color.startswith("#"):
        color = color[1:]

    if len(color) != 6:
        return None

    try:
        int(color, 16)

    except Exception:
        return None

    return f"#{color.upper()}"


def _extract_template_palette():
    if not TEMPLATE_PATH.exists():
        return []

    color_counts = {}
    source_lookup = {}

    try:
        with zipfile.ZipFile(TEMPLATE_PATH, "r") as package:
            for name in package.namelist():
                if not name.endswith(".xml"):
                    continue

                if not (
                    name.startswith("ppt/theme/")
                    or name.startswith("ppt/slideMasters/")
                    or name.startswith("ppt/slideLayouts/")
                    or name.startswith("ppt/slides/")
                ):
                    continue

                text = package.read(name).decode("utf-8", errors="ignore")

                matches = re.findall(
                    r"<a:srgbClr[^>]*\sval=\"([0-9A-Fa-f]{6})\"",
                    text
                )

                matches += re.findall(
                    r"\s(?:color|fill|stroke)=\"?#?([0-9A-Fa-f]{6})\"?",
                    text
                )

                for match in matches:
                    color = _clean_hex_color(match)

                    if not color:
                        continue

                    color_counts[color] = color_counts.get(color, 0) + 1
                    source_lookup.setdefault(color, set()).add(name)

    except Exception:
        return []

    return [
        {
            "hex": color,
            "count": color_counts[color],
            "sources": sorted(source_lookup.get(color, []))[:5],
        }
        for color in sorted(
            color_counts,
            key=lambda item: color_counts[item],
            reverse=True
        )
    ]


def _font(size=14):
    try:
        return ImageFont.truetype("arial.ttf", size)
    except Exception:
        return ImageFont.load_default()


def _load_template_presentation():
    if not TEMPLATE_PATH.exists():
        raise HTTPException(
            status_code=404,
            detail=f"Template not found: {TEMPLATE_PATH}"
        )

    converted = tempfile.NamedTemporaryFile(delete=False, suffix=".pptx")
    converted.close()

    with zipfile.ZipFile(TEMPLATE_PATH, "r") as source:
        with zipfile.ZipFile(converted.name, "w", zipfile.ZIP_DEFLATED) as target:
            for item in source.infolist():
                data = source.read(item.filename)

                if item.filename == "[Content_Types].xml":
                    data = (
                        data.decode("utf-8")
                        .replace(POTX_CONTENT_TYPE, PPTX_CONTENT_TYPE)
                        .encode("utf-8")
                    )

                target.writestr(item, data)

    return Presentation(converted.name)


def _find_libreoffice():
    configured_path = os.environ.get("LIBREOFFICE_PATH")

    if configured_path and Path(configured_path).exists():
        return configured_path

    for command in ("soffice", "libreoffice"):
        resolved = shutil.which(command)

        if resolved:
            return resolved

    windows_candidates = [
        Path("C:/Program Files/LibreOffice/program/soffice.exe"),
        Path("C:/Program Files (x86)/LibreOffice/program/soffice.exe"),
    ]

    for candidate in windows_candidates:
        if candidate.exists():
            return str(candidate)

    return None


def _powerpoint_available():
    if platform.system().lower() != "windows":
        return False

    try:
        import win32com.client  # noqa: F401
        return True

    except Exception:
        return False


def _relative_luminance(rgb):
    red, green, blue = rgb[:3]
    return (
        0.2126 * red
        + 0.7152 * green
        + 0.0722 * blue
    )


def _outline_color_for_box(image, box):
    x1, y1, x2, y2 = box
    sample_points = [
        (
            max(x1 + 4, 0),
            max(y1 + 4, 0)
        ),
        (
            min(x2 - 4, image.width - 1),
            max(y1 + 4, 0)
        ),
        (
            max(x1 + 4, 0),
            min(y2 - 4, image.height - 1)
        ),
        (
            min(x2 - 4, image.width - 1),
            min(y2 - 4, image.height - 1)
        ),
        (
            min(max(int((x1 + x2) / 2), 0), image.width - 1),
            min(max(int((y1 + y2) / 2), 0), image.height - 1)
        ),
    ]

    luminance = sum(
        _relative_luminance(
            image.getpixel(point)
        )
        for point in sample_points
    ) / len(sample_points)

    if luminance < 120:
        return "#FFFFFF", "#111827"

    return "#111827", "#FFFFFF"


def _draw_placeholder_outline(draw, box, label, line_color, label_fill):
    x1, y1, x2, y2 = box
    dash = 8
    gap = 5

    for x in range(x1, x2, dash + gap):
        draw.line(
            (
                x,
                y1,
                min(x + dash, x2),
                y1
            ),
            fill=line_color,
            width=2
        )
        draw.line(
            (
                x,
                y2,
                min(x + dash, x2),
                y2
            ),
            fill=line_color,
            width=2
        )

    for y in range(y1, y2, dash + gap):
        draw.line(
            (
                x1,
                y,
                x1,
                min(y + dash, y2)
            ),
            fill=line_color,
            width=2
        )
        draw.line(
            (
                x2,
                y,
                x2,
                min(y + dash, y2)
            ),
            fill=line_color,
            width=2
        )

    if (x2 - x1) > 90 and (y2 - y1) > 28:
        draw.rectangle(
            (
                x1 + 5,
                y1 + 5,
                min(x1 + 135, x2 - 5),
                y1 + 24
            ),
            fill=label_fill
        )
        draw.text(
            (
                x1 + 9,
                y1 + 7
            ),
            label[:22],
            fill=line_color,
            font=_font(10)
        )


def _draw_icon_box(draw, box, line_color, label):
    x1, y1, x2, y2 = box

    draw.rounded_rectangle(
        box,
        radius=3,
        outline=line_color,
        width=1
    )

    cx = int(
        (x1 + x2) / 2
    )
    cy = int(
        (y1 + y2) / 2
    )

    if label == "table":
        for offset in (-5, 0, 5):
            draw.line(
                (
                    x1 + 4,
                    cy + offset,
                    x2 - 4,
                    cy + offset
                ),
                fill=line_color,
                width=1
            )
            draw.line(
                (
                    cx + offset,
                    y1 + 4,
                    cx + offset,
                    y2 - 4
                ),
                fill=line_color,
                width=1
            )

    elif label == "chart":
        for idx, height in enumerate((7, 12, 17)):
            left = x1 + 5 + (idx * 7)
            draw.rectangle(
                (
                    left,
                    y2 - 5 - height,
                    left + 4,
                    y2 - 5
                ),
                fill=line_color
            )

    elif label == "image":
        draw.rectangle(
            (
                x1 + 5,
                y1 + 5,
                x2 - 5,
                y2 - 5
            ),
            outline=line_color,
            width=1
        )
        draw.polygon(
            [
                (
                    x1 + 6,
                    y2 - 6
                ),
                (
                    cx - 2,
                    cy
                ),
                (
                    cx + 3,
                    cy + 5
                ),
                (
                    x2 - 6,
                    y2 - 6
                ),
            ],
            outline=line_color
        )

    elif label == "video":
        draw.rectangle(
            (
                x1 + 5,
                y1 + 6,
                x2 - 5,
                y2 - 6
            ),
            outline=line_color,
            width=1
        )
        draw.polygon(
            [
                (
                    cx - 3,
                    cy - 6
                ),
                (
                    cx - 3,
                    cy + 6
                ),
                (
                    cx + 7,
                    cy
                ),
            ],
            fill=line_color
        )

    else:
        draw.ellipse(
            (
                cx - 7,
                cy - 7,
                cx + 7,
                cy + 7
            ),
            outline=line_color,
            width=1
        )


def _draw_placeholder_insert_icons(draw, box, kind, line_color):
    x1, y1, x2, y2 = box
    width = x2 - x1
    height = y2 - y1

    if width < 110 or height < 70:
        return

    if kind == "Picture":
        icons = ["image"]
    elif kind == "Content":
        icons = [
            "table",
            "chart",
            "image",
            "video",
        ]
    else:
        return

    icon_size = 28
    gap = 7
    total_width = (len(icons) * icon_size) + ((len(icons) - 1) * gap)
    start_x = int(
        x1 + ((width - total_width) / 2)
    )
    start_y = int(
        y1 + ((height - icon_size) / 2)
    )

    for index, icon in enumerate(icons):
        icon_left = start_x + (index * (icon_size + gap))
        _draw_icon_box(
            draw,
            (
                icon_left,
                start_y,
                icon_left + icon_size,
                start_y + icon_size
            ),
            line_color,
            icon
        )


def _overlay_placeholder_outlines(canvas, prs, layout):
    draw = ImageDraw.Draw(
        canvas
    )

    for placeholder in layout.placeholders:
        kind = _placeholder_kind(
            placeholder
        )
        box = _scaled_box(
            placeholder,
            prs.slide_width,
            prs.slide_height
        )
        line_color, label_fill = _outline_color_for_box(
            canvas,
            box
        )
        _draw_placeholder_outline(
            draw,
            box,
            kind,
            line_color,
            label_fill
        )
        _draw_placeholder_insert_icons(
            draw,
            box,
            kind,
            line_color
        )


def _image_file_to_data_url(path, prs=None, layout=None):
    with Image.open(path) as image:
        image.thumbnail(
            (
                PREVIEW_WIDTH,
                PREVIEW_HEIGHT
            )
        )

        canvas = Image.new(
            "RGB",
            (
                PREVIEW_WIDTH,
                PREVIEW_HEIGHT
            ),
            "#F8FAFC"
        )

        left = int(
            (PREVIEW_WIDTH - image.width) / 2
        )
        top = int(
            (PREVIEW_HEIGHT - image.height) / 2
        )

        canvas.paste(
            image.convert("RGB"),
            (
                left,
                top
            )
        )

        if prs is not None and layout is not None:
            _overlay_placeholder_outlines(
                canvas,
                prs,
                layout
            )

        buffer = io.BytesIO()
        canvas.save(
            buffer,
            format="PNG"
        )

    encoded = base64.b64encode(
        buffer.getvalue()
    ).decode("ascii")

    return f"data:image/png;base64,{encoded}"


def _create_layout_preview_deck(prs):
    preview_prs = _load_template_presentation()

    for index in range(len(preview_prs.slide_layouts)):
        preview_prs.slides.add_slide(
            preview_prs.slide_layouts[index]
        )

    temp_file = tempfile.NamedTemporaryFile(
        delete=False,
        suffix=".pptx"
    )

    preview_prs.save(
        temp_file.name
    )

    return temp_file.name


def _render_layout_previews_with_powerpoint(prs):
    global LAST_RENDER_ERROR

    if not _powerpoint_available():
        return None

    try:
        import pythoncom
        import win32com.client

    except Exception:
        return None

    deck_path = str(
        Path(_create_layout_preview_deck(prs)).resolve()
    )

    output_dir = Path(
        tempfile.mkdtemp(
            prefix="layout_poc_powerpoint_previews_"
        )
    )

    powerpoint = None
    presentation = None

    try:
        pythoncom.CoInitialize()

        powerpoint = win32com.client.DispatchEx(
            "PowerPoint.Application"
        )

        powerpoint.Visible = 1

        presentation = powerpoint.Presentations.Open(
            deck_path,
            WithWindow=False
        )

        previews = []

        for index in range(1, presentation.Slides.Count + 1):
            image_path = output_dir / f"layout_{index:03d}.png"

            presentation.Slides(index).Export(
                str(image_path),
                "PNG",
                PREVIEW_WIDTH,
                PREVIEW_HEIGHT
            )

            previews.append(
                _image_file_to_data_url(
                    image_path,
                    prs,
                    prs.slide_layouts[index - 1]
                )
            )

        if len(previews) < len(prs.slide_layouts):
            return None

        return previews[:len(prs.slide_layouts)]

    except Exception as exc:
        LAST_RENDER_ERROR = f"PowerPoint renderer failed: {type(exc).__name__}: {exc}"
        return None

    finally:
        try:
            if presentation:
                presentation.Close()

        except Exception:
            pass

        try:
            if powerpoint:
                powerpoint.Quit()

        except Exception:
            pass

        try:
            pythoncom.CoUninitialize()

        except Exception:
            pass


def _render_layout_previews_with_libreoffice(prs):
    global LAST_RENDER_ERROR

    libreoffice = _find_libreoffice()

    if not libreoffice:
        return None

    deck_path = _create_layout_preview_deck(
        prs
    )

    output_dir = Path(
        tempfile.mkdtemp(
            prefix="layout_poc_previews_"
        )
    )

    try:
        subprocess.run(
            [
                libreoffice,
                "--headless",
                "--convert-to",
                "png",
                "--outdir",
                str(output_dir),
                deck_path,
            ],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=90,
        )

        image_files = sorted(
            list(output_dir.glob("*.png"))
            + list(output_dir.glob("*.PNG"))
        )

        if len(image_files) < len(prs.slide_layouts):
            return None

        previews = [
            _image_file_to_data_url(
                image_files[index],
                prs,
                prs.slide_layouts[index]
            )
            for index in range(len(prs.slide_layouts))
        ]

        return previews

    except Exception as exc:
        LAST_RENDER_ERROR = f"LibreOffice renderer failed: {type(exc).__name__}: {exc}"
        return None


def _render_layout_previews(prs):
    global RENDERED_PREVIEW_CACHE
    global RENDERED_PREVIEW_SOURCE
    global LAST_RENDER_ERROR

    LAST_RENDER_ERROR = None

    source_key = (
        str(TEMPLATE_PATH),
        TEMPLATE_PATH.stat().st_mtime
    )

    if (
        RENDERED_PREVIEW_CACHE is not None
        and RENDERED_PREVIEW_SOURCE == source_key
    ):
        return RENDERED_PREVIEW_CACHE

    renderers = [
        (
            "powerpoint",
            "Microsoft PowerPoint",
            _render_layout_previews_with_powerpoint,
        ),
        (
            "libreoffice",
            _find_libreoffice() or "LibreOffice",
            _render_layout_previews_with_libreoffice,
        ),
    ]

    for mode, renderer, render_func in renderers:
        previews = render_func(
            prs
        )

        if previews:
            RENDERED_PREVIEW_CACHE = {
                "mode": mode,
                "renderer": renderer,
                "previews": previews,
            }
            RENDERED_PREVIEW_SOURCE = source_key
            return RENDERED_PREVIEW_CACHE

    return {
        "mode": "schematic",
        "renderer": None,
        "previews": None,
        "error": LAST_RENDER_ERROR,
    }


def _placeholder_kind(placeholder):
    try:
        placeholder_type = str(placeholder.placeholder_format.type)
    except Exception:
        placeholder_type = ""

    if "TITLE" in placeholder_type:
        return "Title"
    if "PICTURE" in placeholder_type:
        return "Picture"
    if "BODY" in placeholder_type:
        return "Body"
    if "OBJECT" in placeholder_type:
        return "Content"
    return "Placeholder"


def _placeholder_colors(kind):
    colors = {
        "Title": ("#DDEBFF", "#3B73B9"),
        "Body": ("#E9F7EF", "#39895A"),
        "Content": ("#FFF3D6", "#B87800"),
        "Picture": ("#F2E8FF", "#8056B3"),
        "Placeholder": ("#EEF2F7", "#64748B"),
    }
    return colors.get(kind, colors["Placeholder"])


def _scaled_box(shape, slide_width, slide_height):
    left = int(shape.left / slide_width * PREVIEW_WIDTH)
    top = int(shape.top / slide_height * PREVIEW_HEIGHT)
    width = int(shape.width / slide_width * PREVIEW_WIDTH)
    height = int(shape.height / slide_height * PREVIEW_HEIGHT)

    return (
        left,
        top,
        max(left + width, left + 1),
        max(top + height, top + 1),
    )


def _layout_preview_data_url(prs, layout):
    image = Image.new("RGB", (PREVIEW_WIDTH, PREVIEW_HEIGHT), "#F8FAFC")
    draw = ImageDraw.Draw(image)

    draw.rectangle(
        (0, 0, PREVIEW_WIDTH - 1, PREVIEW_HEIGHT - 1),
        outline="#CBD5E1",
        width=2,
    )
    draw.text((16, 14), layout.name, fill="#0F172A", font=_font(18))

    for placeholder in layout.placeholders:
        kind = _placeholder_kind(placeholder)
        fill, outline = _placeholder_colors(kind)
        box = _scaled_box(placeholder, prs.slide_width, prs.slide_height)
        draw.rounded_rectangle(box, radius=8, fill=fill, outline=outline, width=2)
        draw.text(
            (box[0] + 8, box[1] + 7),
            f"{kind}: {placeholder.name}"[:38],
            fill=outline,
            font=_font(12),
        )

    if not layout.placeholders:
        draw.text((24, 76), "No placeholders", fill="#64748B", font=_font(15))

    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    encoded = base64.b64encode(buffer.getvalue()).decode("ascii")
    return f"data:image/png;base64,{encoded}"


def _layout_metadata(prs, layout, index, rendered_preview=None):
    placeholders = []

    for placeholder in layout.placeholders:
        placeholder_format = placeholder.placeholder_format
        placeholders.append({
            "idx": placeholder_format.idx,
            "type": str(placeholder_format.type),
            "kind": _placeholder_kind(placeholder),
            "name": placeholder.name,
            "left": int(placeholder.left),
            "top": int(placeholder.top),
            "width": int(placeholder.width),
            "height": int(placeholder.height),
        })

    return {
        "index": index,
        "name": layout.name,
        "placeholder_count": len(placeholders),
        "placeholders": placeholders,
        "preview": rendered_preview or _layout_preview_data_url(prs, layout),
        "preview_mode": "rendered" if rendered_preview else "schematic",
    }


def _first_placeholder(slide, placeholder_types):
    for shape in slide.placeholders:
        try:
            if shape.placeholder_format.type in placeholder_types:
                return shape
        except Exception:
            continue
    return None


def _add_title(slide, title):
    title_shape = _first_placeholder(slide, {PP_PLACEHOLDER.TITLE})

    if title_shape:
        title_shape.text = title
        return

    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.35), Inches(12.0), Inches(0.6))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.size = Pt(24)


def _content_bounds(slide, prs):
    content_shape = _first_placeholder(
        slide,
        {
            PP_PLACEHOLDER.BODY,
            PP_PLACEHOLDER.OBJECT,
        },
    )

    if content_shape:
        return content_shape.left, content_shape.top, content_shape.width, content_shape.height

    return (
        Inches(0.55),
        Inches(1.15),
        prs.slide_width - Inches(1.1),
        prs.slide_height - Inches(1.7),
    )


def _add_bullets(slide, prs, bullets, has_chart):
    left, top, width, height = _content_bounds(slide, prs)

    if has_chart:
        width = int(width * 0.36)

    box = slide.shapes.add_textbox(left, top, width, height)
    text_frame = box.text_frame
    text_frame.clear()
    text_frame.word_wrap = True

    for index, bullet in enumerate(bullets):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.text = f"- {bullet}"
        paragraph.font.size = Pt(13)
        paragraph.space_after = Pt(7)

    return left, top, width, height


def _add_chart(slide, prs, request, bullet_bounds):
    left, top, width, height = _content_bounds(slide, prs)

    chart_left = left
    chart_top = top
    chart_width = width
    chart_height = height

    if bullet_bounds:
        bullet_left, bullet_top, bullet_width, bullet_height = bullet_bounds
        chart_left = bullet_left + bullet_width + Inches(0.3)
        chart_top = bullet_top
        chart_width = prs.slide_width - chart_left - Inches(0.55)
        chart_height = bullet_height

    chart_data = CategoryChartData()
    chart_data.categories = request.categories
    chart_data.add_series("Sample", request.values)

    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        chart_left,
        chart_top,
        chart_width,
        chart_height,
        chart_data,
    ).chart

    chart.has_title = True
    chart.chart_title.text_frame.text = request.chart_title
    chart.has_legend = False

    selected_colors = [
        _clean_hex_color(color)
        for color in request.selected_colors
    ]
    selected_colors = [
        color
        for color in selected_colors
        if color
    ]

    if selected_colors:
        try:
            series = chart.series[0]

            for index, point in enumerate(series.points):
                color = selected_colors[index % len(selected_colors)][1:]
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = RGBColor.from_string(
                    color
                )

        except Exception:
            pass

    try:
        chart.category_axis.tick_labels.font.size = Pt(8)
        chart.value_axis.tick_labels.font.size = Pt(8)
    except Exception:
        pass


@app.get("/", response_class=HTMLResponse)
def index():
    return (STATIC_DIR / "index.html").read_text(encoding="utf-8")


@app.get("/api/layouts")
def get_layouts():
    prs = _load_template_presentation()
    preview_result = _render_layout_previews(
        prs
    )
    previews = preview_result.get(
        "previews"
    )

    return {
        "template": str(TEMPLATE_PATH),
        "preview_mode": preview_result.get(
            "mode"
        ),
        "renderer": preview_result.get(
            "renderer"
        ),
        "renderer_error": preview_result.get(
            "error"
        ),
        "palette": _extract_template_palette(),
        "slide_width": int(prs.slide_width),
        "slide_height": int(prs.slide_height),
        "layouts": [
            _layout_metadata(
                prs,
                layout,
                index,
                (
                    previews[index]
                    if previews
                    else None
                )
            )
            for index, layout in enumerate(prs.slide_layouts)
        ],
    }


@app.post("/api/render-slide")
def render_slide(request: RenderSlideRequest):
    prs = _load_template_presentation()

    if request.layout_index < 0 or request.layout_index >= len(prs.slide_layouts):
        raise HTTPException(status_code=400, detail="Invalid layout index")

    if len(request.categories) != len(request.values):
        raise HTTPException(
            status_code=400,
            detail="categories and values must have the same length",
        )

    slide = prs.slides.add_slide(prs.slide_layouts[request.layout_index])

    _add_title(slide, request.title)

    bullet_bounds = _add_bullets(
        slide=slide,
        prs=prs,
        bullets=request.bullets,
        has_chart=request.include_chart,
    )

    if request.include_chart:
        _add_chart(slide, prs, request, bullet_bounds)

    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".pptx")
    prs.save(temp_file.name)

    return FileResponse(
        temp_file.name,
        media_type=(
            "application/vnd.openxmlformats-"
            "officedocument.presentationml.presentation"
        ),
        filename="layout_poc_render.pptx",
    )
